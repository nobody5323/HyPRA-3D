"""Word（.docx）解析器（可选依赖 `python-docx`，**惰性导入**）。

老式 `.doc` 是 OLE 复合二进制格式，`python-docx` 读不了——那种文件的拦截
不在这里，而在 `factory.select_parser()`（按扩展名就地拒绝）。
"""

from __future__ import annotations

import io

from app.memory.knowledge.parser.base import DocumentParser, ParseError


class DocxParser(DocumentParser):
    """提取 .docx 的段落与表格文字。"""

    name = "docx"
    suffixes = frozenset({".docx"})
    dependency_hint = "python-docx"

    def available(self) -> bool:
        try:
            import docx  # noqa: F401, PLC0415 - 可选依赖探测
        except ImportError:
            return False
        return True

    def extract(self, data: bytes) -> str:
        from docx import Document  # noqa: PLC0415 - 可选依赖

        try:
            document = Document(io.BytesIO(data))
        except Exception as exc:  # noqa: BLE001
            raise ParseError(f"Word 文档解析失败：{exc}") from exc

        parts = [paragraph.text for paragraph in document.paragraphs]
        # 表格是常见的信息载体（对照表、参数表），不能漏
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells]
                joined = " | ".join(cell for cell in cells if cell)
                if joined:
                    parts.append(joined)

        text = "\n".join(part for part in parts if part.strip())
        if not text.strip():
            raise ParseError("该 Word 文档没有可提取的文字内容")
        return text
