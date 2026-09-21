"""文档解析：把上传的文件转成纯文本（随后交给 cleaner 清洗）。

支持 txt / md / pdf / docx 四种格式。

**两个明确的能力边界**——遇到时**报错**，而不是静默入库空文档：

1. **扫描版 PDF**（图片型）没有文字层，`extract_text()` 返回空。
   这是最常见的一类失败：用户以为传上去了，实际上库里只有空白。
   因此这里主动检测并提示需要先做 OCR。
2. **老式 `.doc`** 是 OLE 复合二进制格式，`python-docx` 读不了。
   直接拦截并提示「另存为 .docx」，比抛一个用户看不懂的异常更有用。

解析结果**不做任何清洗**——职责分离：parser 负责「取出文字」，
cleaner 负责「去噪音」，两者的测试可以各自独立。
"""

from __future__ import annotations

import io
from dataclasses import dataclass
from pathlib import Path

from docx import Document
from pypdf import PdfReader

#: 纯文本类扩展名（含日志、CSV 等实际是文本的格式）
_PLAIN_SUFFIXES = {".txt", ".text", ".log", ".csv", ".tsv"}
#: Markdown 扩展名（分块时标题会作为一级语义边界）
_MARKDOWN_SUFFIXES = {".md", ".markdown", ".mdown"}
#: 文本解码的候选编码（按顺序尝试；utf-8-sig 同时处理 BOM）
_TEXT_ENCODINGS = ("utf-8-sig", "utf-8", "gb18030")

#: 支持的全部扩展名（用于错误提示与 API 层校验）
SUPPORTED_SUFFIXES = frozenset(
    _PLAIN_SUFFIXES | _MARKDOWN_SUFFIXES | {".pdf", ".docx"}
)


class ParseError(ValueError):
    """解析失败（格式不支持 / 文件损坏 / 无可提取文字）。"""


@dataclass
class ParseResult:
    """解析结果。"""

    text: str
    #: 来源格式：text | md | pdf | docx
    source_type: str
    #: 文档标题（默认取文件名主干，可被调用方覆盖）
    title: str


def _decode_text(data: bytes) -> str:
    """按候选编码解码文本（中文常见 GB18030）。"""
    for encoding in _TEXT_ENCODINGS:
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ParseError("无法识别文本编码（支持 UTF-8 与 GB18030）")


def _parse_pdf(data: bytes) -> str:
    """提取 PDF 文字层。"""
    try:
        reader = PdfReader(io.BytesIO(data))
        pages = [page.extract_text() or "" for page in reader.pages]
    except Exception as exc:  # noqa: BLE001 - 第三方库异常类型不稳定
        raise ParseError(f"PDF 解析失败：{exc}") from exc

    text = "\n".join(pages)
    if not text.strip():
        raise ParseError(
            "该 PDF 没有可提取的文字层（可能是扫描件或图片型 PDF）。"
            "请先用 OCR 工具转成文本，或直接上传 .txt / .md"
        )
    return text


def _parse_docx(data: bytes) -> str:
    """提取 .docx 的段落与表格文字。"""
    try:
        document = Document(io.BytesIO(data))
    except Exception as exc:  # noqa: BLE001
        raise ParseError(f"Word 文档解析失败：{exc}") from exc

    parts = [paragraph.text for paragraph in document.paragraphs]
    # 表格是常见的信息载体（对照表、参数表），不能漏
    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            joined = " | ".join(cell for cell in cells if cell)
            if joined:
                parts.append(joined)

    text = "\n".join(part for part in parts if part.strip())
    if not text.strip():
        raise ParseError("该 Word 文档没有可提取的文字内容")
    return text


def _title_from(filename: str) -> str:
    """从文件名推断标题（去扩展名）。"""
    return Path(filename or "").stem.strip() or "未命名文档"


def parse(filename: str, data: bytes, *, title: str = "") -> ParseResult:
    """按扩展名解析上传的文件。

    参数:
        filename: 原始文件名（用于判断格式与推断标题）；
        data: 文件字节；
        title: 显式标题（留空则取文件名主干）。

    抛出:
        ParseError: 格式不支持、文件损坏、或无法提取到文字。
    """
    if not data:
        raise ParseError("文件内容为空")

    suffix = Path(filename or "").suffix.lower()

    if suffix in _PLAIN_SUFFIXES or suffix in _MARKDOWN_SUFFIXES:
        text = _decode_text(data)
        source_type = "md" if suffix in _MARKDOWN_SUFFIXES else "text"
    elif suffix == ".pdf":
        text, source_type = _parse_pdf(data), "pdf"
    elif suffix == ".docx":
        text, source_type = _parse_docx(data), "docx"
    elif suffix == ".doc":
        raise ParseError(
            "不支持老式 .doc 格式（二进制复合文档）："
            "请在 Word 中另存为 .docx 后重新上传"
        )
    else:
        raise ParseError(
            f"不支持的文件类型「{suffix or '无扩展名'}」："
            "支持 .txt / .md / .pdf / .docx"
        )

    if not text.strip():
        raise ParseError("文件没有可提取的文字内容")

    return ParseResult(
        text=text,
        source_type=source_type,
        title=title.strip() or _title_from(filename),
    )
