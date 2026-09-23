"""文档解析测试。

PDF **正向路径**没有覆盖：生成带文字层的 PDF 需要额外依赖（reportlab），
不值得为测试引入。这里重点覆盖 PDF 的**错误路径**——「扫描件被静默当成空文档
入库」才是真正会造成损失的失败模式。docx 与文本格式则正反向都覆盖。
"""

import io

import pytest
from docx import Document
from pypdf import PdfWriter

from app.memory.knowledge.parser import (
    SUPPORTED_SUFFIXES,
    DocxParser,
    MarkdownParser,
    ParseError,
    PdfParser,
    PlainTextParser,
    parse,
    select_parser,
)


def _docx_bytes(paragraphs: list[str], *, table: list[list[str]] | None = None) -> bytes:
    document = Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table:
        rows, cols = len(table), len(table[0])
        doc_table = document.add_table(rows=rows, cols=cols)
        for r, row in enumerate(table):
            for c, cell in enumerate(row):
                doc_table.cell(r, c).text = cell
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _blank_pdf_bytes() -> bytes:
    """生成一个无文字层的 PDF（模拟扫描件）。"""
    writer = PdfWriter()
    writer.add_blank_page(width=200, height=200)
    buffer = io.BytesIO()
    writer.write(buffer)
    return buffer.getvalue()


# ---------- 文本 / Markdown ----------


def test_parse_plain_text() -> None:
    result = parse("日记.txt", "小林今天说了很多。".encode("utf-8"))

    assert result.source_type == "text"
    assert result.title == "日记"
    assert "小林" in result.text


def test_parse_gbk_encoded_text() -> None:
    """国内常见的 GBK/GB18030 编码必须能读（否则用户会以为文件损坏）。"""
    result = parse("笔记.txt", "今天下雨了，有点想家。".encode("gb18030"))

    assert "今天下雨了" in result.text


def test_parse_markdown_marks_source_type() -> None:
    result = parse("章节.md", "# 标题\n\n正文".encode("utf-8"))

    assert result.source_type == "md"
    assert result.text.startswith("# 标题")


def test_explicit_title_overrides_filename() -> None:
    result = parse("a.txt", b"content", title="我的标题")

    assert result.title == "我的标题"


def test_title_falls_back_to_untitled() -> None:
    """文件名为空白时回退到占位标题（而不是空标题）。"""
    # 注意：".txt" 会被 pathlib 当成隐藏文件（suffix 为空），故用 " .txt"
    assert parse(" .txt", b"x").title == "未命名文档"


# ---------- docx ----------


def test_parse_docx_paragraphs_and_tables() -> None:
    """表格是常见信息载体（对照表、参数表），不能漏。"""
    data = _docx_bytes(
        ["第一段内容", "第二段内容"],
        table=[["项目", "说明"], ["记忆", "分层存储"]],
    )
    result = parse("文档.docx", data)

    assert result.source_type == "docx"
    assert "第一段内容" in result.text
    assert "第二段内容" in result.text
    assert "记忆 | 分层存储" in result.text


def test_empty_docx_raises() -> None:
    with pytest.raises(ParseError, match="没有可提取的文字"):
        parse("空.docx", _docx_bytes([]))


def test_corrupted_docx_raises() -> None:
    with pytest.raises(ParseError, match="Word 文档解析失败"):
        parse("坏.docx", b"not a real docx")


# ---------- pdf ----------


def test_scanned_pdf_without_text_layer_raises() -> None:
    """关键：扫描件没有文字层，必须报错而不是静默入库空文档。"""
    with pytest.raises(ParseError, match="没有可提取的文字层"):
        parse("扫描件.pdf", _blank_pdf_bytes())


def test_corrupted_pdf_raises() -> None:
    with pytest.raises(ParseError, match="PDF 解析失败"):
        parse("坏.pdf", b"%PDF-1.4 broken")


# ---------- 边界与拒绝 ----------


def test_legacy_doc_is_rejected_with_actionable_message() -> None:
    """老 .doc 是二进制复合格式，python-docx 读不了——提示要可操作。"""
    with pytest.raises(ParseError, match="另存为 .docx"):
        parse("旧文档.doc", b"\xd0\xcf\x11\xe0")


def test_unsupported_suffix_raises() -> None:
    with pytest.raises(ParseError, match="不支持的文件类型"):
        parse("图片.png", b"\x89PNG")


def test_empty_file_raises() -> None:
    with pytest.raises(ParseError, match="文件内容为空"):
        parse("空.txt", b"")


def test_whitespace_only_text_raises() -> None:
    with pytest.raises(ParseError, match="没有可提取的文字内容"):
        parse("空白.txt", "   \n\n  ".encode("utf-8"))


# =============================================================
# 解析器选择（拆分新增：一种格式一个实现）
# =============================================================


def test_select_parser_picks_implementation_by_suffix() -> None:
    assert isinstance(select_parser("a.txt"), PlainTextParser)
    assert isinstance(select_parser("a.LOG"), PlainTextParser)      # 扩展名大小写不敏感
    assert isinstance(select_parser("a.md"), MarkdownParser)
    assert isinstance(select_parser("a.markdown"), MarkdownParser)
    assert isinstance(select_parser("a.pdf"), PdfParser)
    assert isinstance(select_parser("a.docx"), DocxParser)


def test_supported_suffixes_cover_every_parser() -> None:
    """`SUPPORTED_SUFFIXES` 由实现聚合而来，不会与实现脱节。"""
    for parser in (PlainTextParser(), MarkdownParser(), PdfParser(), DocxParser()):
        assert parser.suffixes <= SUPPORTED_SUFFIXES, parser.name
        assert parser.name in {"text", "md", "pdf", "docx"}


def test_missing_optional_dependency_is_actionable(monkeypatch) -> None:
    """缺可选依赖时只影响该格式，并给出可执行的安装提示。

    惰性导入是这次拆分顺带修掉的：原来顶层 `from pypdf import ...`，
    少装一个包会让整个 parser 模块 import 失败——连 .txt 上传也起不来。
    """
    monkeypatch.setattr(PdfParser, "available", lambda self: False)

    with pytest.raises(ParseError, match="pip install pypdf"):
        parse("文档.pdf", b"%PDF-1.4")
    # 其它格式不受影响
    assert parse("导语.txt", "你好".encode("utf-8")).source_type == "text"


def test_source_type_distinguishes_text_and_markdown() -> None:
    """两者处理一样但 source_type 必须不同——chunker 靠它决定分块策略。"""
    assert parse("a.txt", b"plain").source_type == "text"
    assert parse("a.md", b"# heading").source_type == "md"
